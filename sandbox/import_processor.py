"""Trusted metadata extractor for uploaded engineering files.

The Worker invokes this module with a fixed command and an allowlisted format.
It never executes macros, repository scripts, Xacro, firmware, or model plugins.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import tarfile
import zipfile
from pathlib import Path, PurePosixPath


MAX_TEXT = 50_000
MAX_ARCHIVE_ENTRIES = 10_000


def safe_archive_path(value: str) -> bool:
    normalized = value.replace("\\", "/")
    path = PurePosixPath(normalized)
    return bool(normalized) and not path.is_absolute() and ".." not in path.parts and "\x00" not in normalized


def archive_metadata(path: Path, format_key: str) -> dict:
    entries: list[dict] = []
    total = 0
    unsafe = 0
    if format_key in {"zip", "usdz"}:
        with zipfile.ZipFile(path) as archive:
            for info in archive.infolist()[:MAX_ARCHIVE_ENTRIES]:
                if not safe_archive_path(info.filename):
                    unsafe += 1
                    continue
                total += info.file_size
                entries.append({"path": info.filename[:1024], "sizeBytes": info.file_size, "compressedBytes": info.compress_size})
    else:
        with tarfile.open(path, mode="r:*") as archive:
            for info in archive.getmembers()[:MAX_ARCHIVE_ENTRIES]:
                if not safe_archive_path(info.name):
                    unsafe += 1
                    continue
                total += max(0, info.size)
                entries.append({"path": info.name[:1024], "sizeBytes": max(0, info.size), "kind": "directory" if info.isdir() else "file"})
    warnings = []
    if unsafe:
        warnings.append(f"Ignored {unsafe} unsafe archive path(s).")
    if len(entries) >= MAX_ARCHIVE_ENTRIES:
        warnings.append(f"Archive inventory was limited to {MAX_ARCHIVE_ENTRIES} entries.")
    return {"metadata": {"entries": entries, "entryCount": len(entries), "uncompressedBytes": total}, "warnings": warnings}


def xlsx_metadata(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    sheets = []
    for sheet in workbook.worksheets[:100]:
        rows = []
        for row in sheet.iter_rows(min_row=1, max_row=51, values_only=True):
            rows.append([None if value is None else str(value)[:500] for value in row[:100]])
        sheets.append({"name": sheet.title[:200], "maxRow": sheet.max_row, "maxColumn": sheet.max_column, "previewRows": rows})
    workbook.close()
    return {"metadata": {"sheets": sheets, "sheetCount": len(sheets)}, "warnings": []}


def pdf_metadata(path: Path) -> dict:
    completed = subprocess.run(
        ["pdftotext", "-layout", "-nopgbrk", str(path), "-"],
        capture_output=True,
        text=True,
        timeout=45,
        check=False,
    )
    text = completed.stdout[:MAX_TEXT]
    warnings = []
    if completed.returncode != 0:
        warnings.append("PDF text extraction failed; the file remains available for manual review.")
    elif not text.strip():
        warnings.append("No machine-readable text was found; the PDF may require OCR.")
    return {"metadata": {"textPreview": text, "textTruncated": len(completed.stdout) > MAX_TEXT}, "warnings": warnings}


def mesh_metadata(path: Path) -> dict:
    completed = subprocess.run(["assimp", "info", str(path)], capture_output=True, text=True, timeout=45, check=False)
    output = (completed.stdout + "\n" + completed.stderr)[:MAX_TEXT]
    metadata: dict = {"assimpSummary": output, "parsed": completed.returncode == 0}
    for key, pattern in {
        "vertices": r"Vertices:\s*(\d+)",
        "faces": r"Faces:\s*(\d+)",
        "meshes": r"Meshes:\s*(\d+)",
        "materials": r"Materials:\s*(\d+)",
        "animations": r"Animations:\s*(\d+)",
    }.items():
        match = re.search(pattern, output, re.IGNORECASE)
        if match:
            metadata[key] = int(match.group(1))
    warnings = [] if completed.returncode == 0 else ["The native mesh parser could not read this file; the original is preserved."]
    return {"metadata": metadata, "warnings": warnings}


def cad_metadata(path: Path, format_key: str) -> dict:
    raw = path.read_bytes()[:8_000_000]
    text = raw.decode("latin-1", errors="ignore")
    metadata = {"byteLengthInspected": len(raw), "format": format_key}
    warnings = ["Geometric CAD conversion is unavailable in this processor; entity metadata was extracted without executing file content."]
    if format_key in {"step", "stp"}:
        metadata["iso10303Entities"] = len(re.findall(r"^#\d+\s*=", text, re.MULTILINE))
        metadata["productNames"] = list(dict.fromkeys(re.findall(r"PRODUCT\s*\(\s*'([^']*)'", text, re.IGNORECASE)))[:100]
    elif format_key in {"iges", "igs"}:
        metadata["directoryEntries"] = sum(1 for line in text.splitlines() if len(line) >= 73 and line[72:73] == "D")
    elif format_key == "usd":
        metadata["primitives"] = len(re.findall(r"\bdef\s+(?:Xform|Mesh|Scope|Material|Physics\w*)\b", text))
    return {"metadata": metadata, "warnings": warnings}


def process(path: Path, format_key: str) -> dict:
    if format_key in {"zip", "tar", "tgz", "tar_gz", "usdz"}:
        result = archive_metadata(path, format_key)
    elif format_key == "xlsx":
        result = xlsx_metadata(path)
    elif format_key == "pdf":
        result = pdf_metadata(path)
    elif format_key in {"stl", "obj", "gltf", "glb", "dae"}:
        result = mesh_metadata(path)
    elif format_key in {"step", "stp", "iges", "igs", "usd"}:
        result = cad_metadata(path, format_key)
    else:
        raise ValueError("Unsupported allowlisted format")
    return {
        "format": format_key,
        "processor": "robopartpicker-sandbox-static-v1",
        "factInference": {"facts": ["Metadata was extracted from the uploaded file."], "inferences": []},
        "missingDependencies": [],
        **result,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", required=True)
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    result = process(Path(args.input), args.format)
    Path(args.output).write_text(json.dumps(result, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main()
