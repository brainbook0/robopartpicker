from __future__ import annotations

import csv
import io
import json
import math
import re
import zipfile
from dataclasses import dataclass
from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import urlsplit
from xml.etree.ElementTree import ParseError

import yaml
from defusedxml.ElementTree import fromstring as safe_xml_fromstring
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_TEXT_BYTES = 8 * 1024 * 1024
MAX_PDF_BYTES = 25 * 1024 * 1024
MAX_METADATA_BYTES = 50 * 1024 * 1024
MAX_NODES = 50_000
MAX_DEPTH = 32
MAX_XML_DEPTH = 64
MAX_LINES = 100_000
MAX_ROWS = 5_000
MAX_COLUMNS = 100
MAX_CELLS = 100_000
MAX_SHEETS = 10
MAX_PDF_PAGES = 100
MAX_PDF_TEXT_CHARS = 2_000_000
MAX_XLSX_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_XLSX_ENTRIES = 1_000

ROBOT_SUFFIXES = {".urdf", ".xacro", ".sdf", ".srdf", ".mjcf"}
NATIVE_METADATA_SUFFIXES = {
    ".dae",
    ".dxf",
    ".f3d",
    ".glb",
    ".gltf",
    ".iges",
    ".igs",
    ".obj",
    ".sldasm",
    ".sldprt",
    ".step",
    ".stl",
    ".stp",
    ".usd",
    ".usdz",
}
ARCHIVE_SUFFIXES = {".7z", ".gz", ".rar", ".tar", ".tgz", ".zip"}


class ParserError(ValueError):
    pass


class ParserLimitExceeded(ParserError):
    pass


class UnsafeDocument(ParserError):
    pass


@dataclass(frozen=True)
class EvidenceSpan:
    locator: str
    value: Any


@dataclass(frozen=True)
class RobotDescription:
    format: str
    name: str | None
    units: dict[str, str]
    links: list[dict[str, Any]]
    joints: list[dict[str, Any]]
    mesh_references: list[str]
    missing_assets: list[str]


@dataclass(frozen=True)
class UnsupportedFormat:
    file_name: str
    file_type: str
    state: str
    byte_size: int
    content_sha256: str
    reason: str


@dataclass(frozen=True)
class ParseResult:
    format: str
    evidence_spans: tuple[EvidenceSpan, ...] = ()
    structured_data: Any = None
    robot_description: RobotDescription | None = None
    unsupported_format: UnsupportedFormat | None = None
    network_access: bool = False


def parse_document(file_name: str, content: bytes) -> ParseResult:
    suffix = Path(file_name).suffix.lower()
    if suffix in ARCHIVE_SUFFIXES:
        raise UnsafeDocument("archive inputs are not parsed or expanded")
    if suffix in NATIVE_METADATA_SUFFIXES:
        _enforce_size(content, MAX_METADATA_BYTES)
        return ParseResult(
            format=suffix.removeprefix("."),
            unsupported_format=UnsupportedFormat(
                file_name=file_name,
                file_type=suffix.removeprefix("."),
                state="metadata_only",
                byte_size=len(content),
                content_sha256=sha256(content).hexdigest(),
                reason="native CAD and mesh content is never executed",
            ),
        )
    if suffix == ".pdf":
        _enforce_size(content, MAX_PDF_BYTES)
        return _parse_pdf(content)

    _enforce_size(content, MAX_TEXT_BYTES)
    if suffix in ROBOT_SUFFIXES:
        return _parse_robot(suffix.removeprefix("."), content)
    if suffix in {".html", ".htm"}:
        return _parse_html(content, suffix.removeprefix("."))
    if suffix in {".md", ".markdown"}:
        return _parse_markdown(content, suffix.removeprefix("."))
    if suffix in {".csv", ".tsv"}:
        return _parse_delimited(content, suffix.removeprefix("."))
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    if suffix == ".json":
        return _parse_json(content)
    if suffix in {".yaml", ".yml"}:
        return _parse_yaml(content, suffix.removeprefix("."))
    if suffix == ".xml":
        return _parse_xml(content)
    return ParseResult(
        format=suffix.removeprefix(".") or "unknown",
        unsupported_format=UnsupportedFormat(
            file_name=file_name,
            file_type=suffix.removeprefix(".") or "unknown",
            state="rejected",
            byte_size=len(content),
            content_sha256=sha256(content).hexdigest(),
            reason="unsupported format",
        ),
    )


def _parse_html(content: bytes, format_name: str) -> ParseResult:
    text = _decode_utf8(content)
    parser = _BoundedHtmlParser()
    try:
        parser.feed(text)
        parser.close()
    except ParserError:
        raise
    except Exception as error:
        raise UnsafeDocument("HTML could not be parsed safely") from error
    return ParseResult(format=format_name, evidence_spans=tuple(parser.spans))


class _BoundedHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.spans: list[EvidenceSpan] = []
        self._nodes = 0
        self._depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self._nodes += 1
        self._depth += 1
        if self._nodes > MAX_NODES:
            raise ParserLimitExceeded("HTML node limit exceeded")
        if self._depth > MAX_DEPTH:
            raise ParserLimitExceeded("HTML depth limit exceeded")

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self._nodes += 1
        if self._nodes > MAX_NODES:
            raise ParserLimitExceeded("HTML node limit exceeded")

    def handle_endtag(self, tag: str) -> None:
        self._depth = max(0, self._depth - 1)

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.spans.append(EvidenceSpan(
                locator=f"html:line={self.getpos()[0]}",
                value=value,
            ))


def _parse_markdown(content: bytes, format_name: str) -> ParseResult:
    lines = _decode_utf8(content).splitlines()
    if len(lines) > MAX_LINES:
        raise ParserLimitExceeded("Markdown line limit exceeded")
    spans = tuple(
        EvidenceSpan(locator=f"markdown:line={line_number}", value=line.strip())
        for line_number, line in enumerate(lines, start=1)
        if line.strip()
    )
    return ParseResult(format=format_name, evidence_spans=spans)


def _parse_delimited(content: bytes, format_name: str) -> ParseResult:
    delimiter = "," if format_name == "csv" else "\t"
    try:
        reader = csv.reader(io.StringIO(_decode_utf8(content)), delimiter=delimiter)
        rows = list(reader)
    except csv.Error as error:
        raise UnsafeDocument("delimited text could not be parsed safely") from error
    if not rows:
        return ParseResult(format=format_name)
    if len(rows) > MAX_ROWS + 1:
        raise ParserLimitExceeded("tabular row limit exceeded")
    headers = rows[0]
    if len(headers) > MAX_COLUMNS:
        raise ParserLimitExceeded("tabular column limit exceeded")
    cells = sum(len(row) for row in rows)
    if cells > MAX_CELLS:
        raise ParserLimitExceeded("tabular cell limit exceeded")
    spans = []
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) != len(headers):
            raise UnsafeDocument("tabular row width does not match its header")
        for column, value in zip(headers, row, strict=True):
            _reject_formula(value)
            if value != "":
                spans.append(EvidenceSpan(
                    locator=f"{format_name}:row={row_number}:column={column}",
                    value=value,
                ))
    return ParseResult(
        format=format_name,
        evidence_spans=tuple(spans),
        structured_data={
            "columns": headers,
            "rowCount": max(0, len(rows) - 1),
        },
    )


def _parse_xlsx(content: bytes) -> ParseResult:
    _preflight_xlsx(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise UnsafeDocument("XLSX could not be parsed safely") from error
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ParserLimitExceeded("XLSX sheet limit exceeded")
        spans = []
        total_cells = 0
        for sheet in workbook.worksheets:
            if sheet.max_row > MAX_ROWS:
                raise ParserLimitExceeded("XLSX row limit exceeded")
            if sheet.max_column > MAX_COLUMNS:
                raise ParserLimitExceeded("XLSX column limit exceeded")
            for row in sheet.iter_rows():
                total_cells += len(row)
                if total_cells > MAX_CELLS:
                    raise ParserLimitExceeded("XLSX cell limit exceeded")
                for cell in row:
                    if cell.data_type == "f":
                        raise UnsafeDocument("XLSX formula cells are not allowed")
                    _reject_formula(cell.value)
                    if cell.value not in {None, ""}:
                        spans.append(EvidenceSpan(
                            locator=f"xlsx:sheet={sheet.title}:cell={cell.coordinate}",
                            value=cell.value,
                        ))
        return ParseResult(
            format="xlsx",
            evidence_spans=tuple(spans),
            structured_data={"sheetCount": len(workbook.worksheets)},
        )
    finally:
        workbook.close()


def _preflight_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise ParserLimitExceeded("XLSX archive entry limit exceeded")
            expanded = sum(entry.file_size for entry in entries)
            if expanded > MAX_XLSX_EXPANDED_BYTES:
                raise ParserLimitExceeded("XLSX expansion limit exceeded")
    except zipfile.BadZipFile as error:
        raise UnsafeDocument("XLSX container is invalid") from error


def _parse_json(content: bytes) -> ParseResult:
    try:
        value = json.loads(_decode_utf8(content))
    except json.JSONDecodeError as error:
        raise UnsafeDocument("JSON could not be parsed safely") from error
    spans = _structured_spans(value, prefix="json")
    return ParseResult(format="json", evidence_spans=tuple(spans), structured_data=value)


def _parse_yaml(content: bytes, format_name: str) -> ParseResult:
    try:
        value = yaml.safe_load(_decode_utf8(content))
    except yaml.YAMLError as error:
        raise UnsafeDocument("YAML could not be parsed safely") from error
    spans = _structured_spans(value, prefix="yaml")
    return ParseResult(format=format_name, evidence_spans=tuple(spans), structured_data=value)


def _structured_spans(value: Any, *, prefix: str) -> list[EvidenceSpan]:
    spans: list[EvidenceSpan] = []
    node_count = 0
    active: set[int] = set()

    def walk(current: Any, pointer: str, depth: int) -> None:
        nonlocal node_count
        node_count += 1
        if node_count > MAX_NODES:
            raise ParserLimitExceeded(f"{prefix.upper()} node limit exceeded")
        if depth > MAX_DEPTH:
            raise ParserLimitExceeded(f"{prefix.upper()} depth limit exceeded")
        if isinstance(current, (dict, list)):
            identity = id(current)
            if identity in active:
                raise UnsafeDocument(f"{prefix.upper()} cyclic aliases are not allowed")
            active.add(identity)
            try:
                if isinstance(current, dict):
                    for key, child in current.items():
                        escaped = str(key).replace("~", "~0").replace("/", "~1")
                        walk(child, f"{pointer}/{escaped}", depth + 1)
                else:
                    for index, child in enumerate(current):
                        walk(child, f"{pointer}/{index}", depth + 1)
            finally:
                active.remove(identity)
            return
        if isinstance(current, float) and not math.isfinite(current):
            raise UnsafeDocument(f"{prefix.upper()} numbers must be finite")
        spans.append(EvidenceSpan(locator=f"{prefix}:{pointer or '/'}", value=current))

    walk(value, "", 0)
    return spans


def _parse_xml(content: bytes) -> ParseResult:
    root = _safe_xml_root(content)
    spans = _xml_spans(root)
    return ParseResult(format="xml", evidence_spans=tuple(spans))


def _safe_xml_root(content: bytes) -> Any:
    try:
        return safe_xml_fromstring(
            content,
            forbid_dtd=True,
            forbid_entities=True,
            forbid_external=True,
        )
    except (DefusedXmlException, ParseError) as error:
        raise UnsafeDocument("XML active content or invalid structure was rejected") from error


def _xml_spans(root: Any) -> list[EvidenceSpan]:
    spans: list[EvidenceSpan] = []
    nodes = 0

    def walk(element: Any, path: str, depth: int) -> None:
        nonlocal nodes
        nodes += 1
        if nodes > MAX_NODES:
            raise ParserLimitExceeded("XML node limit exceeded")
        if depth > MAX_XML_DEPTH:
            raise ParserLimitExceeded("XML depth limit exceeded")
        text = (element.text or "").strip()
        if text:
            spans.append(EvidenceSpan(locator=f"xml:{path}", value=text))
        for name, value in sorted(element.attrib.items()):
            spans.append(EvidenceSpan(locator=f"xml:{path}/@{name}", value=value))
        counts: dict[str, int] = {}
        for child in element:
            tag = _local_name(child.tag)
            counts[tag] = counts.get(tag, 0) + 1
            walk(child, f"{path}/{tag}[{counts[tag]}]", depth + 1)

    walk(root, f"/{_local_name(root.tag)}", 0)
    return spans


def _parse_pdf(content: bytes) -> ParseResult:
    try:
        reader = PdfReader(io.BytesIO(content), strict=True)
    except Exception as error:
        raise UnsafeDocument("PDF could not be parsed safely") from error
    if reader.is_encrypted:
        raise UnsafeDocument("encrypted PDF files are not allowed")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ParserLimitExceeded("PDF page limit exceeded")
    spans = []
    total_chars = 0
    try:
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            total_chars += len(text)
            if total_chars > MAX_PDF_TEXT_CHARS:
                raise ParserLimitExceeded("PDF text limit exceeded")
            if text:
                spans.append(EvidenceSpan(locator=f"pdf:page={page_number}", value=text))
    except ParserError:
        raise
    except Exception as error:
        raise UnsafeDocument("PDF text extraction failed safely") from error
    return ParseResult(
        format="pdf",
        evidence_spans=tuple(spans),
        structured_data={"pageCount": len(reader.pages)},
    )


def _parse_robot(format_name: str, content: bytes) -> ParseResult:
    root = _safe_xml_root(content)
    _xml_spans(root)
    if format_name in {"urdf", "xacro"}:
        robot = _parse_urdf_like(root, format_name)
    elif format_name == "sdf":
        robot = _parse_sdf(root)
    elif format_name == "srdf":
        robot = _parse_srdf(root)
    else:
        robot = _parse_mjcf(root)
    spans = tuple(
        [
            *(EvidenceSpan(f"robot:link={link['name']}", link["name"]) for link in robot.links),
            *(EvidenceSpan(f"robot:joint={joint['name']}", joint["name"]) for joint in robot.joints),
        ],
    )
    return ParseResult(
        format=format_name,
        evidence_spans=spans,
        robot_description=robot,
    )


def _parse_urdf_like(root: Any, format_name: str) -> RobotDescription:
    if _local_name(root.tag) != "robot":
        raise UnsafeDocument(f"{format_name.upper()} root must be robot")
    _validate_xacro_includes(root)
    links = [_urdf_link(element) for element in root if _local_name(element.tag) == "link"]
    joints = [_urdf_joint(element) for element in root if _local_name(element.tag) == "joint"]
    meshes = _mesh_references(root)
    return RobotDescription(
        format=format_name,
        name=root.attrib.get("name"),
        units=_si_units(),
        links=links,
        joints=joints,
        mesh_references=meshes,
        missing_assets=list(meshes),
    )


def _urdf_link(element: Any) -> dict[str, Any]:
    link: dict[str, Any] = {"name": _required_attribute(element, "name", "link")}
    inertial = _child(element, "inertial")
    if inertial is not None:
        mass = _child(inertial, "mass")
        inertia = _child(inertial, "inertia")
        if mass is not None and "value" in mass.attrib:
            link["mass"] = _finite_float(mass.attrib["value"], "link mass")
        if inertia is not None:
            link["inertia"] = {
                key: _finite_float(inertia.attrib[key], f"inertia {key}")
                for key in ("ixx", "ixy", "ixz", "iyy", "iyz", "izz")
                if key in inertia.attrib
            }
    return link


def _urdf_joint(element: Any) -> dict[str, Any]:
    parent = _child(element, "parent")
    child = _child(element, "child")
    if parent is None or child is None:
        raise UnsafeDocument("URDF joint parent and child are required")
    axis = _child(element, "axis")
    limit = _child(element, "limit")
    joint = {
        "name": _required_attribute(element, "name", "joint"),
        "type": element.attrib.get("type", "unknown"),
        "parent": _required_attribute(parent, "link", "joint parent"),
        "child": _required_attribute(child, "link", "joint child"),
        "axis": _vector(axis.attrib.get("xyz", "1 0 0")) if axis is not None else [1.0, 0.0, 0.0],
        "limits": {},
    }
    if limit is not None:
        joint["limits"] = {
            key: _finite_float(limit.attrib[key], f"joint limit {key}")
            for key in ("lower", "upper", "effort", "velocity")
            if key in limit.attrib
        }
    return joint


def _parse_sdf(root: Any) -> RobotDescription:
    if _local_name(root.tag) != "sdf":
        raise UnsafeDocument("SDF root must be sdf")
    model = next((item for item in root.iter() if _local_name(item.tag) == "model"), None)
    if model is None:
        raise UnsafeDocument("SDF model is required")
    links = []
    joints = []
    for element in model:
        tag = _local_name(element.tag)
        if tag == "link":
            link: dict[str, Any] = {"name": _required_attribute(element, "name", "link")}
            mass = next((item for item in element.iter() if _local_name(item.tag) == "mass"), None)
            if mass is not None and (mass.text or "").strip():
                link["mass"] = _finite_float((mass.text or "").strip(), "link mass")
            links.append(link)
        elif tag == "joint":
            parent = _child_text(element, "parent")
            child = _child_text(element, "child")
            joints.append({
                "name": _required_attribute(element, "name", "joint"),
                "type": element.attrib.get("type", "unknown"),
                "parent": parent,
                "child": child,
                "axis": _sdf_axis(element),
                "limits": _sdf_limits(element),
            })
    meshes = _mesh_references(root)
    return RobotDescription(
        format="sdf",
        name=model.attrib.get("name"),
        units=_si_units(),
        links=links,
        joints=joints,
        mesh_references=meshes,
        missing_assets=list(meshes),
    )


def _parse_srdf(root: Any) -> RobotDescription:
    if _local_name(root.tag) != "robot":
        raise UnsafeDocument("SRDF root must be robot")
    links = [
        {"name": element.attrib["name"]}
        for element in root.iter()
        if _local_name(element.tag) == "link" and "name" in element.attrib
    ]
    joints = [
        {
            "name": element.attrib["name"],
            "type": "semantic_reference",
            "parent": None,
            "child": None,
            "axis": [],
            "limits": {},
        }
        for element in root.iter()
        if _local_name(element.tag) == "joint" and "name" in element.attrib
    ]
    return RobotDescription(
        format="srdf",
        name=root.attrib.get("name"),
        units=_si_units(),
        links=links,
        joints=joints,
        mesh_references=[],
        missing_assets=[],
    )


def _parse_mjcf(root: Any) -> RobotDescription:
    if _local_name(root.tag) != "mujoco":
        raise UnsafeDocument("MJCF root must be mujoco")
    links: list[dict[str, Any]] = []
    joints: list[dict[str, Any]] = []

    def walk_body(element: Any, parent: str | None) -> None:
        name = element.attrib.get("name")
        if name:
            links.append({"name": name})
        for child in element:
            tag = _local_name(child.tag)
            if tag == "joint" and "name" in child.attrib:
                range_values = child.attrib.get("range")
                limits: dict[str, float] = {}
                if range_values:
                    values = _vector(range_values, expected=None)
                    if len(values) == 2:
                        limits = {"lower": values[0], "upper": values[1]}
                joints.append({
                    "name": child.attrib["name"],
                    "type": child.attrib.get("type", "hinge"),
                    "parent": parent,
                    "child": name,
                    "axis": _vector(child.attrib.get("axis", "0 0 1")),
                    "limits": limits,
                })
            elif tag == "body":
                walk_body(child, name or parent)

    for element in root.iter():
        if _local_name(element.tag) == "worldbody":
            for child in element:
                if _local_name(child.tag) == "body":
                    walk_body(child, "world")
            break
    meshes = _mjcf_mesh_references(root)
    return RobotDescription(
        format="mjcf",
        name=root.attrib.get("model"),
        units=_si_units(),
        links=links,
        joints=joints,
        mesh_references=meshes,
        missing_assets=list(meshes),
    )


def _mesh_references(root: Any) -> list[str]:
    references = []
    for element in root.iter():
        tag = _local_name(element.tag)
        value = None
        if tag == "mesh":
            value = element.attrib.get("filename")
        elif tag == "uri":
            value = (element.text or "").strip() or None
        if value is not None:
            _validate_asset_reference(value)
            references.append(value)
    return sorted(set(references))


def _mjcf_mesh_references(root: Any) -> list[str]:
    references = []
    for element in root.iter():
        if _local_name(element.tag) == "mesh" and "file" in element.attrib:
            value = element.attrib["file"]
            _validate_asset_reference(value)
            references.append(value)
    return sorted(set(references))


def _validate_xacro_includes(root: Any) -> None:
    for element in root.iter():
        if _local_name(element.tag) == "include" and "filename" in element.attrib:
            _validate_asset_reference(element.attrib["filename"])


def _validate_asset_reference(value: str) -> None:
    if "\\" in value or "\x00" in value:
        raise UnsafeDocument("unsafe asset reference was rejected")
    parsed = urlsplit(value)
    if parsed.scheme:
        if parsed.scheme not in {"package", "model"}:
            raise UnsafeDocument("unsafe asset reference was rejected")
        if not parsed.netloc or parsed.query or parsed.fragment:
            raise UnsafeDocument("unsafe asset reference was rejected")
        path = f"{parsed.netloc}/{parsed.path.lstrip('/')}"
    else:
        path = value
    pure = PurePosixPath(path)
    if pure.is_absolute() or ".." in pure.parts:
        raise UnsafeDocument("unsafe asset reference was rejected")


def _sdf_axis(joint: Any) -> list[float]:
    for element in joint.iter():
        if _local_name(element.tag) == "xyz" and (element.text or "").strip():
            return _vector((element.text or "").strip())
    return [1.0, 0.0, 0.0]


def _sdf_limits(joint: Any) -> dict[str, float]:
    result = {}
    for element in joint.iter():
        tag = _local_name(element.tag)
        if tag in {"lower", "upper", "effort", "velocity"} and (element.text or "").strip():
            result[tag] = _finite_float((element.text or "").strip(), f"joint limit {tag}")
    return result


def _child(element: Any, name: str) -> Any | None:
    return next((child for child in element if _local_name(child.tag) == name), None)


def _child_text(element: Any, name: str) -> str:
    child = _child(element, name)
    if child is None or not (child.text or "").strip():
        raise UnsafeDocument(f"SDF joint {name} is required")
    return (child.text or "").strip()


def _required_attribute(element: Any, name: str, label: str) -> str:
    value = element.attrib.get(name)
    if not value:
        raise UnsafeDocument(f"{label} {name} is required")
    return value


def _vector(value: str, expected: int | None = 3) -> list[float]:
    parts = value.split()
    if expected is not None and len(parts) != expected:
        raise UnsafeDocument(f"vector must contain {expected} values")
    return [_finite_float(part, "vector value") for part in parts]


def _finite_float(value: Any, label: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise UnsafeDocument(f"{label} must be numeric") from error
    if not math.isfinite(result):
        raise UnsafeDocument(f"{label} must be finite")
    return result


def _si_units() -> dict[str, str]:
    return {"length": "m", "angle": "rad", "mass": "kg"}


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _reject_formula(value: Any) -> None:
    if not isinstance(value, str):
        return
    stripped = value.lstrip()
    if stripped.startswith(("=", "+", "@")):
        raise UnsafeDocument("tabular formula cells are not allowed")
    if stripped.startswith("-"):
        try:
            float(stripped)
        except ValueError as error:
            raise UnsafeDocument("tabular formula cells are not allowed") from error


def _decode_utf8(content: bytes) -> str:
    try:
        return content.decode("utf8")
    except UnicodeDecodeError as error:
        raise UnsafeDocument("text content must be UTF-8") from error


def _enforce_size(content: bytes, limit: int) -> None:
    if len(content) > limit:
        raise ParserLimitExceeded(f"document byte limit of {limit} exceeded")
