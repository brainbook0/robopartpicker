from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook


MAX_BOM_BYTES = 2 * 1024 * 1024
MAX_SHEETS = 10
MAX_ROWS = 5_000
MAX_COLUMNS = 50
MAX_CELLS = 100_000


class BomParseError(ValueError):
    pass


def parse_bom(content: bytes, *, file_name: str) -> list[dict[str, Any]]:
    if len(content) > MAX_BOM_BYTES:
        raise BomParseError("BOM exceeds the 2 MiB limit")
    suffix = Path(file_name).suffix.lower()
    if suffix in {".csv", ".tsv"}:
        rows = _parse_delimited(content, delimiter="," if suffix == ".csv" else "\t")
        locator = "csv" if suffix == ".csv" else "tsv"
    elif suffix in {".json", ".yaml", ".yml"}:
        rows = _parse_structured(content, suffix)
        locator = suffix.removeprefix(".")
    elif suffix in {".md", ".markdown"}:
        rows = _parse_markdown(content)
        locator = "markdown"
    elif suffix == ".xlsx":
        return _parse_xlsx(content)
    else:
        raise BomParseError(f"unsupported BOM format {suffix or '<none>'}")
    return _normalize_rows(rows, locator=locator)


def _parse_delimited(content: bytes, *, delimiter: str) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf8")
    except UnicodeDecodeError as error:
        raise BomParseError("BOM text must be UTF-8") from error
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if reader.fieldnames is None:
        raise BomParseError("BOM header is required")
    _check_dimensions(len(reader.fieldnames), 0)
    rows = []
    for row_number, row in enumerate(reader, start=2):
        if row_number > MAX_ROWS + 1:
            raise BomParseError("BOM row limit exceeded")
        rows.append({**row, "_line": row_number})
    return rows


def _parse_structured(content: bytes, suffix: str) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf8")
        value = json.loads(text) if suffix == ".json" else yaml.safe_load(text)
    except (UnicodeDecodeError, json.JSONDecodeError, yaml.YAMLError) as error:
        raise BomParseError("BOM structured content is invalid") from error
    if isinstance(value, dict):
        value = value.get("items", value.get("bom"))
    if not isinstance(value, list) or not all(isinstance(item, dict) for item in value):
        raise BomParseError("BOM must contain an array of item objects")
    if len(value) > MAX_ROWS:
        raise BomParseError("BOM row limit exceeded")
    return [{**item, "_line": index} for index, item in enumerate(value, start=1)]


def _parse_markdown(content: bytes) -> list[dict[str, Any]]:
    try:
        lines = content.decode("utf8").splitlines()
    except UnicodeDecodeError as error:
        raise BomParseError("BOM text must be UTF-8") from error
    table_lines = [line.strip() for line in lines if line.strip().startswith("|")]
    if len(table_lines) < 2:
        raise BomParseError("Markdown BOM table is missing")
    headers = [cell.strip() for cell in table_lines[0].strip("|").split("|")]
    separator = [cell.strip() for cell in table_lines[1].strip("|").split("|")]
    if len(headers) != len(separator) or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in separator):
        raise BomParseError("Markdown BOM table separator is invalid")
    _check_dimensions(len(headers), 0)
    rows = []
    for line_number, line in enumerate(table_lines[2:], start=3):
        values = [cell.strip() for cell in line.strip("|").split("|")]
        if len(values) != len(headers):
            raise BomParseError("Markdown BOM row width does not match its header")
        rows.append({**dict(zip(headers, values, strict=True)), "_line": line_number})
    return rows


def _parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise BomParseError("XLSX BOM could not be opened safely") from error
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise BomParseError("BOM sheet limit exceeded")
        normalized: list[dict[str, Any]] = []
        total_cells = 0
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows()
            try:
                header_cells = next(iterator)
            except StopIteration:
                continue
            if len(header_cells) > MAX_COLUMNS:
                raise BomParseError("BOM column limit exceeded")
            headers = [_cell_value(cell) for cell in header_cells]
            for row_number, cells in enumerate(iterator, start=2):
                if row_number > MAX_ROWS + 1:
                    raise BomParseError("BOM row limit exceeded")
                total_cells += len(cells)
                if total_cells > MAX_CELLS:
                    raise BomParseError("BOM cell limit exceeded")
                row = {
                    str(header): _cell_value(cell)
                    for header, cell in zip(headers, cells, strict=False)
                    if header not in {None, ""}
                }
                row["_line"] = row_number
                row["_sheet"] = sheet.title
                normalized.extend(_normalize_rows(
                    [row],
                    locator=f"xlsx:sheet={sheet.title}",
                ))
        return normalized
    finally:
        workbook.close()


def _cell_value(cell: Any) -> Any:
    if cell.data_type == "f":
        raise BomParseError("BOM cell formulas are not allowed")
    value = cell.value
    _reject_formula(value)
    return value


def _normalize_rows(rows: list[dict[str, Any]], *, locator: str) -> list[dict[str, Any]]:
    normalized = []
    for index, row in enumerate(rows, start=1):
        _check_dimensions(len(row), index)
        for value in row.values():
            _reject_formula(value)
        line = int(row.get("_line", index))
        part_name = _pick(row, "part_name", "part name", "part", "name")
        if part_name in {None, ""}:
            raise BomParseError(f"BOM row {line} is missing a part name")
        quantity = _number(_pick(row, "quantity", "qty"), field="quantity", integer=True)
        unit_price = _number(_pick(row, "unit_price", "unit price", "price"), field="unit_price")
        currency_value = _pick(row, "currency")
        currency = str(currency_value or "").upper() or None
        if currency is not None and not re.fullmatch(r"[A-Z]{3}", currency):
            raise BomParseError(f"BOM row {line} has an invalid currency")
        sheet = row.get("_sheet")
        evidence_locator = (
            f"{locator}:row={line}"
            if not str(locator).startswith("xlsx:")
            else f"{locator}:row={line}"
        )
        normalized.append({
            "lineId": str(line),
            "partName": str(part_name),
            "manufacturerPartNumber": _optional_text(
                _pick(
                    row,
                    "manufacturer_part_number",
                    "manufacturer part number",
                    "mpn",
                    "part_number",
                ),
            ),
            "quantity": quantity,
            "unitPrice": unit_price,
            "currency": currency,
            "evidenceLocator": evidence_locator,
        })
    return normalized


def _pick(row: dict[str, Any], *names: str) -> Any:
    lowered = {
        str(key).strip().lower().replace("-", "_"): value
        for key, value in row.items()
        if not str(key).startswith("_")
    }
    for name in names:
        candidate = name.lower().replace("-", "_")
        if candidate in lowered:
            return lowered[candidate]
    return None


def _number(value: Any, *, field: str, integer: bool = False) -> int | float | None:
    if value in {None, ""}:
        return None
    try:
        parsed = float(value)
    except (TypeError, ValueError) as error:
        raise BomParseError(f"BOM {field} must be numeric") from error
    if parsed < 0:
        raise BomParseError(f"BOM {field} cannot be negative")
    if integer:
        if not parsed.is_integer():
            raise BomParseError("BOM quantity must be a whole number")
        return int(parsed)
    return parsed


def _optional_text(value: Any) -> str | None:
    if value in {None, ""}:
        return None
    return str(value)


def _reject_formula(value: Any) -> None:
    if not isinstance(value, str):
        return
    stripped = value.lstrip()
    if stripped.startswith(("=", "+", "@")):
        raise BomParseError("BOM cell formulas are not allowed")
    if stripped.startswith("-"):
        try:
            float(stripped)
        except ValueError as error:
            raise BomParseError("BOM cell formulas are not allowed") from error


def _check_dimensions(columns: int, row: int) -> None:
    if columns > MAX_COLUMNS:
        raise BomParseError("BOM column limit exceeded")
    if row * max(columns, 1) > MAX_CELLS:
        raise BomParseError("BOM cell limit exceeded")
